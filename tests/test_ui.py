from datetime import date
from decimal import Decimal
import time
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QLockFile, QMimeData, QPoint, QPointF, QSettings, Qt, QTimer, QUrl
from PySide6.QtGui import QDragEnterEvent, QDropEvent
from PySide6.QtNetwork import QLocalServer
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QFileDialog, QMessageBox

from app import MainWindow, MISSING_COLOR, notify_running_instance, serve_instance_requests
from rebate_engine import load_csv


def wait_idle(window, qapp, timeout=10):
    deadline = time.monotonic() + timeout
    while window.worker is not None and time.monotonic() < deadline:
        qapp.processEvents()
        QTest.qWait(5)
    assert window.worker is None, "Import thread did not finish"
    qapp.processEvents()


def test_second_launch_reuses_existing_instance(qapp, tmp_path):
    first = QLockFile(str(tmp_path / "instance.lock"))
    second = QLockFile(str(tmp_path / "instance.lock"))
    assert first.tryLock(0)
    assert not second.tryLock(0)
    server = QLocalServer()
    if not server.listen(f"rebate-test-{uuid4().hex}"):
        first.unlock()
        pytest.skip(f"local sockets unavailable: {server.errorString()}")
    calls = []

    class ExistingWindow:
        def showNormal(self):
            calls.append("show")

        def raise_(self):
            calls.append("raise")

        def activateWindow(self):
            calls.append("activate")

        def import_files(self, paths):
            calls.append(paths)

    server.newConnection.connect(lambda: serve_instance_requests(server, ExistingWindow()))
    try:
        notify_running_instance(server.serverName(), ["a.csv", "b.csv"])
        deadline = time.monotonic() + 1
        while len(calls) < 4 and time.monotonic() < deadline:
            qapp.processEvents()
            QTest.qWait(5)
        assert calls == ["show", "raise", "activate", ["a.csv", "b.csv"]]
    finally:
        server.close()
        first.unlock()


@pytest.fixture
def window(qapp, tmp_path):
    win = MainWindow(settings=QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat))
    win.show()
    qapp.processEvents()
    yield win
    if win.worker:
        win.cancel_import()
        wait_idle(win, qapp)
    win.close()
    qapp.processEvents()


def test_initial_empty_ui(window):
    assert window.total_label.text() == "—"
    assert not window.query_button.isEnabled()
    assert not window.export_button.isEnabled()
    assert window.start_date.isEnabled()
    assert window.start_date.text() == ""


def test_picker_import_and_reimport_replace(window, make_csv, qapp, monkeypatch):
    path = make_csv()
    monkeypatch.setattr(QFileDialog, "getOpenFileNames", lambda *args: ([str(path)], "CSV"))
    window.drop_area.choose_button.click()
    wait_idle(window, qapp)
    assert window.result.total == Decimal("3.65")
    assert window.total_label.text() == "3.65000000"
    assert "MB · 最新 2026-08-25" in window.file_label.text().splitlines()[0]
    window.import_file(str(path))
    wait_idle(window, qapp)
    assert window.result.count == 8
    assert window.result.total == Decimal("3.65")
    other = make_csv([["spot", "8", "0", "9", "2026-08-19"]], filename="other.csv")
    window.import_file(str(other))
    wait_idle(window, qapp)
    assert window.result.count == 1
    assert window.result.total == 9
    assert "001" not in window.dataset.ids


def test_filters_keep_old_results_until_calculate(window, make_csv, qapp):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    window.id_input.setPlainText("001，2,001")
    assert window.result.total == Decimal("3.65")
    assert window.total_label.text() == "3.65000000"
    assert "上次统计结果" in window.scope_label.text()
    assert not window.export_button.isEnabled()
    QTest.keyClicks(window.start_date, "20260818")
    QTest.keyClicks(window.end_date, "20260824")
    assert window.start_date.text() == "2026-08-18"
    assert window.end_date.text() == "2026-08-24"
    window.query_button.click()
    assert window.result.total == Decimal("0.4")
    assert window.result.count == 3
    window.copy_button.click()
    assert qapp.clipboard().text() == "0.40000000"
    window.reset_button.click()
    assert window.result.total == Decimal("0.4")
    assert "上次统计结果" in window.scope_label.text()
    assert window.start_date.text() == ""
    assert window.end_date.text() == ""
    window.query_button.click()
    assert window.result.total == Decimal("3.65")
    assert "上次统计结果" not in window.scope_label.text()


def test_manual_date_range_counts_all_friends_and_inclusive_days(window, make_csv, qapp):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    QTest.keyClicks(window.start_date, "20260818")
    QTest.keyClicks(window.end_date, "20260824")
    assert window.result.total == Decimal("3.65")
    window.query_button.click()
    assert window.result.count == 6
    assert window.result.total == Decimal("2.35")
    assert window.result.filters.ids == ()
    assert "所有好友" in window.scope_label.text()

    window.start_date.clear()
    assert window.result.total == Decimal("2.35")
    window.query_button.click()
    assert window.result.count == 7
    assert window.result.total == Decimal("3.35")


def test_date_digits_insert_separators_without_picker(window):
    QTest.keyClicks(window.start_date, "20260920")
    assert window.start_date.text() == "2026-09-20"
    QTest.keyClick(window.start_date, Qt.Key.Key_Backspace)
    QTest.keyClick(window.start_date, Qt.Key.Key_Backspace)
    assert window.start_date.text() == "2026-09"
    QTest.keyClicks(window.start_date, "30")
    assert window.start_date.text() == "2026-09-30"


@pytest.mark.parametrize("value", ["2026/08/18", "2026-02-30", "2026-8-18", "2026-08-18 12:00:00"])
def test_invalid_manual_date_keeps_old_result_without_stale_export(window, make_csv, qapp, value):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    window.start_date.setText(value)
    window.query_button.click()
    assert window.result.total == Decimal("3.65")
    assert window.total_label.text() == "3.65000000"
    assert not window.export_button.isEnabled()
    assert "开始日期" in window.warning_label.text()


def test_reversed_dates_no_stale_export(window, make_csv, qapp):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    window.start_date.setText("2026-08-25")
    window.end_date.setText("2026-08-18")
    window.calculate()
    assert window.result.total == Decimal("3.65")
    assert "开始日期" in window.warning_label.text()
    assert not window.export_button.isEnabled()


def test_date_range_past_csv_coverage_shows_warning(window, make_csv, qapp):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    window.end_date.setText("2026-08-30")
    window.calculate()
    assert window.result.total == Decimal("3.65")
    assert "文件最新返佣记录仅到 2026-08-25" in window.warning_label.text()
    assert window.warning_label.isVisible()


def test_pending_filters_keep_footer_and_clear_error_tooltip(window, make_csv, qapp):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    assert "本次匹配 8 条" in window.table_footer.text()
    window.id_input.setPlainText("999")
    assert "本次匹配 8 条" in window.table_footer.text()
    assert "上次统计结果" in window.scope_label.text()
    window.warning_label.setToolTip("旧的 UID 名单")
    window.show_error("新的错误")
    assert window.warning_label.toolTip() == ""
    window.invalidate_result()
    assert window.warning_label.text() == ""


def test_real_drag_drop_event(window, make_csv, qapp):
    path = make_csv()
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(path))])
    enter = QDragEnterEvent(QPoint(10, 10), Qt.DropAction.CopyAction, mime, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    qapp.sendEvent(window.drop_area, enter)
    assert enter.isAccepted()
    drop = QDropEvent(QPointF(10, 10), Qt.DropAction.CopyAction, mime, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    qapp.sendEvent(window.drop_area, drop)
    assert drop.isAccepted()
    wait_idle(window, qapp)
    assert window.result.count == 8


def test_multiple_file_drop_accepted(window, make_csv, qapp):
    first, second = make_csv(), make_csv(filename="other.csv")
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(first)), QUrl.fromLocalFile(str(second))])
    event = QDropEvent(QPointF(10, 10), Qt.DropAction.CopyAction, mime, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    window.drop_area.dropEvent(event)
    assert event.isAccepted()
    wait_idle(window, qapp)
    assert len(window.dataset.sources) == 2
    assert window.result.count == 16
    assert window.result.total == Decimal("7.30")


def test_cancel_never_publishes_partial_dataset(window, make_csv, qapp):
    path = make_csv([["spot", "1", "0", "1", "2026-08-18"]] * 20000)
    window.import_file(str(path))
    window.cancel_import()
    wait_idle(window, qapp)
    assert window.dataset is None
    assert window.result is None
    assert not window.export_button.isEnabled()


def test_import_error_not_partial(window, make_csv, qapp):
    path = make_csv([["spot", "1", "0", "1", "2026-08-18"], ["spot", "1", "0", "bad", "2026-08-18"]])
    window.import_file(str(path))
    wait_idle(window, qapp)
    assert window.dataset is None
    assert window.result is None
    assert "第 3 行" in window.last_error
    assert window.drop_area.isEnabled()


def test_export_button_uses_current_tab_and_filters(window, make_csv, tmp_path, qapp, monkeypatch):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    window.id_input.setPlainText("001")
    window.calculate()
    window.tabs.setCurrentIndex(1)
    target = tmp_path / "export.csv"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(target), "CSV"))
    window.export_button.click()
    text = target.read_text(encoding="utf-8-sig")
    assert "0.20000000" in text
    assert "2026-08-18" in text
    assert "3.65000000" not in text


def test_ui_can_export_xlsx_with_filter_selected_extension(window, make_csv, tmp_path, qapp, monkeypatch):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    window.id_input.setPlainText("001")
    window.calculate()
    target_without_suffix = tmp_path / "exact-export"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: (
        str(target_without_suffix), "Excel 工作簿 (*.xlsx)"))
    window.export_button.click()
    target = target_without_suffix.with_suffix(".xlsx")
    assert target.exists()
    workbook = load_workbook(target, data_only=False)
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    uid_column = headers.index("好友ID（现货）") + 1
    amount_column = headers.index("返佣收入(USDT)") + 1
    assert worksheet.cell(2, uid_column).value == "001"
    assert worksheet.cell(2, uid_column).data_type == "s"
    assert worksheet.cell(2, amount_column).value == "0.20000000"
    assert worksheet.cell(2, amount_column).data_type == "s"
    workbook.close()


def test_excel_validation_error_is_visible_and_keeps_result(window, make_csv, tmp_path, qapp, monkeypatch):
    target = tmp_path / "invalid.xlsx"
    # An unknown order label is legal CSV input but may be invalid Excel text.
    path = make_csv([["custom\x01type", "001", "0", "1", "2026-08-18"]], filename="special.csv")
    window.import_file(str(path))
    wait_idle(window, qapp)
    window.order_combo.setCurrentIndex(window.order_combo.findData("custom\x01type"))
    window.calculate()
    total = window.result.total
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(target), "Excel 工作簿 (*.xlsx)"))
    window.export_current()
    assert "字符" in window.last_error and "CSV" in window.last_error
    assert window.warning_label.isVisible()
    assert window.result.total == total
    assert not target.exists()


def test_table_sorts_amount_numerically(window, make_csv, qapp):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    model = window.models[0]
    model.sort(2, Qt.SortOrder.DescendingOrder)
    assert model.rows[0][0] == "123"
    assert model.rows[-1][0] == "3"


def test_close_during_import_waits_for_thread(window, make_csv, qapp):
    path = make_csv([["spot", "1", "0", "1", "2026-08-18"]] * 20000)
    window.import_file(str(path))
    window.close()
    wait_idle(window, qapp)
    assert not window.isVisible()
    assert window.dataset is None


@pytest.mark.parametrize("approve", [True, False])
def test_overwrite_prompt_after_extension_normalization(window, make_csv, tmp_path, qapp, monkeypatch, approve):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    target = tmp_path / "existing.csv"
    target.write_text("keep", encoding="utf8")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(target.with_suffix("")), "CSV"))
    confirmations = []
    def confirm(*args):
        confirmations.append(args[2])
        return QMessageBox.StandardButton.Yes if approve else QMessageBox.StandardButton.No
    monkeypatch.setattr(QMessageBox, "question", confirm)
    window.export_current()
    assert len(confirmations) == 1
    assert "existing.csv" in confirmations[0]
    assert (target.read_text(encoding="utf-8-sig") != "keep") == approve


def test_ui_event_loop_responds_during_large_import(window, make_csv, qapp):
    path = make_csv([["spot", "1", "0", "0.00000001", "2026-08-18"]] * 100000)
    ticks = []
    timer = QTimer()
    timer.setInterval(2)
    timer.timeout.connect(lambda: ticks.append(window.worker is not None))
    window.import_file(str(path))
    timer.start()
    wait_idle(window, qapp)
    timer.stop()
    assert any(ticks), "UI event loop did not respond while import was active"
    assert window.result.count == 100000
    assert window.result.total == Decimal("0.00100000")


def test_multi_picker_and_batch_replacement(window, make_csv, qapp, monkeypatch):
    a, b = make_csv(filename="a.csv"), make_csv(filename="b.csv")
    monkeypatch.setattr(QFileDialog, "getOpenFileNames", lambda *args: ([str(a), str(b)], "CSV"))
    window.choose_file()
    wait_idle(window, qapp)
    assert window.result.count == 16
    assert "2 个 CSV" in window.file_label.text()
    c = make_csv([["spot", "999", "0", "7", "2026-08-18"]], filename="c.csv")
    window.import_files([c])
    wait_idle(window, qapp)
    assert window.result.total == 7
    assert window.dataset.ids == frozenset({"999"})


def test_remembers_last_successful_batch_and_picker_directory(window, make_csv, qapp, monkeypatch, tmp_path):
    a = make_csv(filename="a.csv")
    b = make_csv(filename="b.csv")
    window.import_files([a, b])
    wait_idle(window, qapp)
    settings_file = window.settings.fileName()
    reopened = MainWindow(settings=QSettings(settings_file, QSettings.Format.IniFormat))
    reopened.show()
    try:
        reopened.restore_last_files()
        wait_idle(reopened, qapp)
        assert {source.path for source in reopened.dataset.sources} == {a.resolve(), b.resolve()}
        assert reopened.result.count == 16
        seen = []
        replacement = make_csv([["spot", "9", "0", "4", "2026-08-19"]], filename="new.csv")
        def pick(*args):
            seen.append(args[2])
            return [str(replacement)], "CSV"
        monkeypatch.setattr(QFileDialog, "getOpenFileNames", pick)
        reopened.choose_file()
        wait_idle(reopened, qapp)
        assert seen == [str(tmp_path)]
        assert reopened.result.total == 4
        saved = QSettings(settings_file, QSettings.Format.IniFormat)
        assert saved.value("last_files") == [str(replacement.resolve())]
    finally:
        reopened.close()


def test_failed_replacement_keeps_last_successful_file(window, make_csv, qapp):
    good = make_csv()
    window.import_file(good)
    wait_idle(window, qapp)
    previous = window.settings.value("last_files")
    bad = make_csv([["spot", "8", "0", "bad", "2026-08-19"]], filename="bad.csv")
    window.import_file(bad)
    wait_idle(window, qapp)
    assert window.dataset is None
    assert window.settings.value("last_files") == previous


def test_restore_uses_saved_encoding(window, make_csv, qapp):
    path = make_csv(encoding="gb18030")
    window.encoding_combo.setCurrentIndex(window.encoding_combo.findData("gb18030"))
    window.import_file(path)
    wait_idle(window, qapp)
    reopened = MainWindow(settings=QSettings(window.settings.fileName(), QSettings.Format.IniFormat))
    reopened.show()
    try:
        reopened.restore_last_files()
        wait_idle(reopened, qapp)
        assert reopened.encoding_combo.currentData() == "gb18030"
        assert reopened.result.total == Decimal("3.65")
    finally:
        reopened.close()


def test_complete_26_uid_list_status_search_and_copy(window, make_csv, qapp):
    ids = [str(10000000 + i) for i in range(26)]
    path = make_csv([["spot", uid, "0", "1", "2026-08-17"] for uid in ids])
    window.import_file(str(path))
    wait_idle(window, qapp)
    window.id_input.setPlainText("\n".join(ids + ["99999999"]))
    window.start_date.setText("2026-08-18")
    window.calculate()
    assert len(window.models[0].rows) == 27
    assert "未找到 1 个" in window.audit_label.text()
    assert "当前无记录 26" in window.audit_label.text()
    assert not window.warning_label.isVisible()
    window.status_filter.setCurrentIndex(window.status_filter.findData("inactive"))
    assert len(window.models[0].rows) == 26
    assert window.result.total == 0
    window.copy_ids_button.click()
    assert qapp.clipboard().text().splitlines() == ids
    window.uid_search.setText("10000025")
    assert len(window.models[0].rows) == 1
    assert window.models[0].rows[0][0] == "10000025"
    window.id_input.clear()
    assert window.result is not None and not window.copy_ids_button.isEnabled()
    assert "上次统计结果" in window.scope_label.text()


def test_status_filter_keeps_settlement_total_and_export_matches_view(window, make_csv, tmp_path, qapp, monkeypatch):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    total = window.total_label.text()
    window.status_filter.setCurrentIndex(window.status_filter.findData("zero"))
    assert window.total_label.text() == total
    assert len(window.models[0].rows) == 1
    target = tmp_path / "zero.csv"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(target), "CSV"))
    window.export_current()
    import csv
    with target.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 1 and rows[0]["好友ID（现货）"] == "4"
    window.tabs.setCurrentIndex(1)
    assert not window.uid_controls.isVisible()


def test_summary_counts_toggle_uid_table_filter(window, make_csv, qapp):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    window.id_input.setPlainText("001 2 12 123 3 4")
    window.start_date.setText("2026-08-18")
    window.end_date.setText("2026-08-24")
    window.calculate()
    total = window.result.total
    assert window.audit_label.isVisible()
    for status, expected in (("positive", {"001", "2", "123"}),
                             ("negative", {"3"}), ("zero", {"4"}),
                             ("inactive", {"12"}), ("missing", set())):
        assert f'href="{status}"' in window.audit_label.text()
        window.uid_search.setText("001")
        window.tabs.setCurrentIndex(1)
        window.audit_label.linkActivated.emit(status)
        assert window.tabs.currentIndex() == 0
        assert window.uid_search.text() == ""
        assert window.status_filter.currentData() == status
        assert {row[0] for row in window.models[0].rows} == expected
        assert window.result.total == total
        window.audit_label.linkActivated.emit(status)
        assert window.status_filter.currentData() == ""
        assert len(window.models[0].rows) == 6


def test_missing_count_link_toggles_missing_uid_rows(window, make_csv, qapp):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    window.id_input.setPlainText("001 999")
    window.calculate()
    assert window.missing_panel.isVisible()
    assert 'href="missing"' in window.missing_label.text()
    assert 'href="positive"' in window.missing_label.text()
    window.missing_label.linkActivated.emit("missing")
    assert window.status_filter.currentData() == "missing"
    assert [row[0] for row in window.models[0].rows] == ["999"]
    window.missing_label.linkActivated.emit("missing")
    assert window.status_filter.currentData() == ""
    assert {row[0] for row in window.models[0].rows} == {"001", "999"}
    window.missing_label.linkActivated.emit("positive")
    assert window.status_filter.currentData() == "positive"
    assert [row[0] for row in window.models[0].rows] == ["001"]


def test_bad_second_file_clears_previous_batch(window, make_csv, qapp):
    a = make_csv(filename="a.csv")
    bad = make_csv([["spot", "1", "0", "bad", "2026-08-18"]], filename="bad.csv")
    window.import_file(str(a))
    wait_idle(window, qapp)
    window.import_files([a, bad])
    wait_idle(window, qapp)
    assert window.dataset is None and window.result is None
    assert "bad.csv" in window.last_error
    assert not window.file_details_button.isEnabled()
    assert "本次匹配" not in window.table_footer.text()


def test_small_window_primary_actions_visible(window, qapp):
    window.resize(960, 680)
    qapp.processEvents()
    for widget in (window.query_button, window.reset_button, window.export_button, window.total_label):
        assert widget.isVisible()
        assert window.rect().contains(widget.mapTo(window, widget.rect().center()))


def test_cancel_on_second_file_publishes_no_batch(window, make_csv, qapp):
    rows = [["spot", "1", "0", "1E-8", "2026-08-18"]] * 10000
    a, b = make_csv(rows, filename="a.csv"), make_csv(rows, filename="b.csv")
    window.import_files([a, b])
    window.worker.file_started.connect(lambda index, total, name: window.cancel_import() if index == 2 else None)
    wait_idle(window, qapp)
    assert window.dataset is None and window.result is None
    assert not window.export_button.isEnabled()


def test_missing_previous_order_does_not_silently_broaden_filter(window, make_csv, qapp):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    window.order_combo.setCurrentIndex(window.order_combo.findData("USDT-futures"))
    window.calculate()
    path = make_csv([["spot", "1", "0", "3", "2026-08-18"]], filename="spot.csv")
    window.import_file(str(path))
    wait_idle(window, qapp)
    assert window.order_combo.currentData() == "USDT-futures"
    assert window.result.total == 0
    assert window.result.statuses == {"1": "inactive"}


def test_known_order_type_variants_share_one_ui_filter(window, make_csv, qapp):
    rows = [
        ["Spot", "1", "0", "1", "2026-08-18"],
        ["SPOT", "1", "0", "2", "2026-08-18"],
        ["USDT-Futures", "1", "0", "4", "2026-08-18"],
        ["usdt-futures", "1", "0", "8", "2026-08-18"],
    ]
    window.import_file(str(make_csv(rows)))
    wait_idle(window, qapp)
    assert [window.order_combo.itemData(i) for i in range(window.order_combo.count())] == [
        None, "USDT-futures", "spot"
    ]
    window.order_combo.setCurrentIndex(window.order_combo.findData("spot"))
    window.calculate()
    assert window.result.total == 3
    window.order_combo.setCurrentIndex(window.order_combo.findData("USDT-futures"))
    window.calculate()
    assert window.result.total == 12


def test_missing_ids_have_complete_red_record_and_copy(window, make_csv, qapp):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    missing = [str(9000000000 + i) for i in range(26)]
    window.id_input.setPlainText("\n".join(["001", "12", "4", *missing, missing[0]]))
    window.start_date.setText("2026-08-18")
    window.calculate()
    assert window.missing_panel.isVisible()
    assert "未找到 26 个" in window.missing_label.text()
    assert window.missing_list.toPlainText().split("、") == missing
    window.resize(960, 680)
    qapp.processEvents()
    assert window.tabs.widget(0).viewport().height() >= 64
    assert window.missing_list.isReadOnly()
    window.copy_missing_button.click()
    assert qapp.clipboard().text().splitlines() == missing
    model = window.models[0]
    for row_index, row in enumerate(model.rows):
        if row[0] in missing:
            for column in range(model.columnCount()):
                assert model.data(model.index(row_index, column), Qt.ItemDataRole.ForegroundRole).name() == MISSING_COLOR
        elif row[0] in ("12", "4"):
            assert model.data(model.index(row_index, 0), Qt.ItemDataRole.ForegroundRole) is None
    # A table-only search/status filter must not hide IDs from the dedicated record.
    window.uid_search.setText("001")
    window.tabs.setCurrentIndex(1)
    total = window.result.total
    window.show_missing_button.click()
    assert window.tabs.currentIndex() == 0
    assert window.uid_search.text() == ""
    assert window.status_filter.currentData() == "missing"
    assert len(window.models[0].rows) == 26
    assert window.result.total == total
    assert window.missing_list.toPlainText().split("、") == missing


def test_missing_record_stays_until_recalculated_or_batch_replaced(window, make_csv, qapp):
    window.import_file(str(make_csv()))
    wait_idle(window, qapp)
    window.id_input.setPlainText("999")
    window.calculate()
    assert window.missing_panel.isVisible()
    window.id_input.setPlainText("001")
    assert window.missing_panel.isVisible()
    assert not window.copy_missing_button.isEnabled()
    assert window.missing_list.toPlainText() == "999"
    assert "上次统计结果" in window.scope_label.text()
    window.calculate()
    assert "未找到 0 个" in window.audit_label.text()
    assert window.audit_label.isVisible()
    assert not window.missing_panel.isVisible()
    window.id_input.setPlainText("999")
    window.calculate()
    second = make_csv([["spot", "999", "0", "0", "2026-08-18"]], filename="second.csv")
    window.import_file(str(second))
    wait_idle(window, qapp)
    assert not window.missing_panel.isVisible()
    assert window.missing_list.toPlainText() == ""
    assert window.result.statuses["999"] == "zero"
    assert window.models[0].data(window.models[0].index(0, 0), Qt.ItemDataRole.ForegroundRole) is None
