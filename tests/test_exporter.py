import csv
import os
import subprocess
import sys
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from exporter import EXPORT_HEADERS, export_csv, export_result, export_xlsx
from rebate_engine import DataError, Filters, load_csv, load_csvs


@pytest.mark.parametrize("view", ["id", "day", "daily_id"])
def test_export_matches_exact_filtered_total(make_csv, tmp_path, view):
    dataset = load_csv(make_csv())
    result = dataset.query(Filters(("001", "2", "999"), date(2026, 8, 18), date(2026, 8, 24)))
    path = tmp_path / f"out-{view}.csv"
    count = export_csv(path, dataset, result, view)
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert count == len(rows)
    assert sum(Decimal(row["返佣收入(USDT)"]) for row in rows) == result.total
    assert sum(int(row["返佣笔数"]) for row in rows) == result.count
    assert rows[0]["好友ID筛选"] == "001 2 999"
    assert all(row["好友ID筛选"] == "" for row in rows[1:])
    assert all(row["开始日期(含)"] == "2026-08-18" for row in rows)
    if view != "day":
        assert any(row["好友ID（现货）"] == "001" for row in rows)
    if view != "id":
        assert all(row["UID状态"] == "" and row["全批次笔数"] == "" for row in rows)


@pytest.mark.parametrize("view", ["id", "day", "daily_id"])
def test_xlsx_export_preserves_exact_text_and_filtered_total(make_csv, tmp_path, view):
    dataset = load_csv(make_csv())
    result = dataset.query(Filters(("001", "2", "999"), date(2026, 8, 18), date(2026, 8, 24)))
    path = tmp_path / f"out-{view}.xlsx"
    count = export_xlsx(path, dataset, result, view)
    workbook = load_workbook(path, data_only=False)
    worksheet = workbook.active
    values = list(worksheet.values)
    assert list(values[0]) == EXPORT_HEADERS
    rows = [dict(zip(EXPORT_HEADERS, row)) for row in values[1:]]
    assert count == len(rows)
    assert sum(Decimal(row["返佣收入(USDT)"]) for row in rows) == result.total
    assert sum(int(row["返佣笔数"]) for row in rows) == result.count
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == f"A1:N{count + 1}"
    amount_column = EXPORT_HEADERS.index("返佣收入(USDT)") + 1
    for row in range(2, count + 2):
        assert worksheet.cell(row, amount_column).data_type == "s"
        assert worksheet.cell(row, amount_column).number_format == "@"
    if view != "day":
        uid_column = EXPORT_HEADERS.index("好友ID（现货）") + 1
        uid_cell = next(worksheet.cell(row, uid_column) for row in range(2, count + 2)
                        if worksheet.cell(row, uid_column).value == "001")
        assert uid_cell.data_type == "s" and uid_cell.number_format == "@"
    workbook.close()


def test_export_result_dispatches_and_rejects_unknown_suffix(make_csv, tmp_path):
    dataset = load_csv(make_csv())
    result = dataset.query()
    assert export_result(tmp_path / "out.csv", dataset, result) == len(result.per_id)
    assert export_result(tmp_path / "out.xlsx", dataset, result) == len(result.per_id)
    with pytest.raises(DataError, match="CSV 或 XLSX"):
        export_result(tmp_path / "out.txt", dataset, result)


@pytest.mark.parametrize("writer", [export_csv, export_xlsx])
@pytest.mark.parametrize("kind", ["original", "symlink", "hardlink"])
def test_source_overwrite_protection(make_csv, tmp_path, kind, writer):
    source = make_csv()
    original = source.read_bytes()
    target = source
    if kind != "original":
        target = tmp_path / "alias.csv"
        if kind == "symlink":
            target.symlink_to(source)
        else:
            os.link(source, target)
    dataset = load_csv(source)
    with pytest.raises(DataError, match="覆盖源"):
        writer(target, dataset, dataset.query())
    assert source.read_bytes() == original


@pytest.mark.parametrize("writer,suffix", [(export_csv, ".csv"), (export_xlsx, ".xlsx")])
def test_failed_export_preserves_existing_file(make_csv, tmp_path, monkeypatch, writer, suffix):
    dataset = load_csv(make_csv())
    target = tmp_path / f"existing{suffix}"
    target.write_text("existing", encoding="utf8")
    def fail(*args):
        raise OSError("test failure")
    monkeypatch.setattr("exporter.os.replace", fail)
    with pytest.raises(OSError):
        writer(target, dataset, dataset.query())
    assert target.read_text() == "existing"
    assert not list(tmp_path.glob(".rebate-export-*"))


def test_formula_injection_in_filename_is_escaped(make_csv, tmp_path):
    dataset = load_csv(make_csv(filename="=1+1.csv"))
    target = tmp_path / "out.csv"
    export_csv(target, dataset, dataset.query())
    with target.open(encoding="utf-8-sig", newline="") as stream:
        row = next(csv.DictReader(stream))
    assert row["来源文件"].startswith("'=")
    xlsx = tmp_path / "out.xlsx"
    export_xlsx(xlsx, dataset, dataset.query())
    workbook = load_workbook(xlsx, data_only=False)
    source_cell = workbook.active.cell(2, 1)
    assert source_cell.value.startswith("'=") and source_cell.data_type == "s"
    workbook.close()


def test_moved_source_export_still_works_and_source_stays_protected(make_csv, tmp_path):
    original = make_csv()
    dataset = load_csv(original)
    moved = original.rename(tmp_path / "moved.csv")
    existing = tmp_path / "existing.csv"
    existing.write_text("old", encoding="utf8")
    assert export_csv(existing, dataset, dataset.query()) == 6
    assert export_csv(tmp_path / "new.csv", dataset, dataset.query()) == 6
    with pytest.raises(DataError, match="覆盖源"):
        export_csv(moved, dataset, dataset.query())


def test_all_batch_sources_and_aliases_protected(make_csv, tmp_path):
    a, b = make_csv(filename="a.csv"), make_csv(filename="b.csv")
    dataset = load_csvs([a, b])
    alias = tmp_path / "b-alias.csv"
    os.link(b, alias)
    for path in (a, b, alias):
        with pytest.raises(DataError, match="覆盖源"):
            export_csv(path, dataset, dataset.query())


def test_large_id_list_written_only_once(make_csv, tmp_path):
    ids = tuple(str(10000000 + i) for i in range(1000))
    dataset = load_csv(make_csv([["spot", uid, "0", "1", "2026-08-18"] for uid in ids]))
    output = tmp_path / "large.csv"
    export_csv(output, dataset, dataset.query(Filters(ids)))
    assert output.stat().st_size < 180_000
    with output.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 1000
    assert sum(bool(row["好友ID筛选"]) for row in rows) == 1
    assert rows[0]["好友ID筛选"].split() == list(ids)


def test_status_export_contains_complete_inactive_missing_and_zero(make_csv, tmp_path):
    dataset = load_csv(make_csv())
    result = dataset.query(Filters(("001", "12", "4", "999"), date(2026, 8, 18)))
    output = tmp_path / "status.csv"
    export_csv(output, dataset, result, status="no_rebate")
    with output.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert {row["好友ID（现货）"] for row in rows} == {"12", "4"}
    assert all(row["所选文件中存在"] == "是" for row in rows)
    assert {row["UID状态"] for row in rows} == {"存在 · 当前条件无记录", "有记录 · 净额为 0"}
    export_csv(output, dataset, result, status="missing")
    with output.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 1 and rows[0]["所选文件中存在"] == "否"
    assert rows[0]["好友ID（现货）"] == "999"


def test_missing_excel_dependency_does_not_disable_app_or_csv(make_csv, tmp_path):
    code = """
import builtins
original_import = builtins.__import__
def without_excel(name, *args, **kwargs):
    if name == 'openpyxl' or name.startswith('openpyxl.'):
        raise ModuleNotFoundError("No module named 'openpyxl'", name='openpyxl')
    return original_import(name, *args, **kwargs)
builtins.__import__ = without_excel
import sys
from pathlib import Path
from PySide6.QtWidgets import QApplication
from app import MainWindow
from exporter import export_csv, export_xlsx
from rebate_engine import DataError, load_csv
application = QApplication([])
window = MainWindow()
window.show()
application.processEvents()
assert window.isVisible()
dataset = load_csv(sys.argv[1])
assert export_csv(sys.argv[2], dataset, dataset.query()) == 6
try:
    export_xlsx(sys.argv[3], dataset, dataset.query())
except DataError as exc:
    assert 'pip install -r requirements.txt' in str(exc)
else:
    raise AssertionError('Missing Excel dependency was not reported')
assert not Path(sys.argv[3]).exists()
window.close()
"""
    result = subprocess.run([sys.executable, "-c", code, str(make_csv()),
                             str(tmp_path / "out.csv"), str(tmp_path / "out.xlsx")],
                            capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stderr


@pytest.mark.parametrize("order", ["custom\x01type", "custom\ufffetype"])
def test_invalid_excel_text_fails_cleanly_and_preserves_existing_file(make_csv, tmp_path, order):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES
    dataset = load_csv(make_csv([[order, "1", "0", "1", "2026-08-18"]]))
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"keep existing export")
    temp_files_before = set(ALL_TEMP_FILES)
    with pytest.raises(DataError, match="字符.*CSV"):
        export_xlsx(target, dataset, dataset.query(Filters(order_type=order)))
    assert target.read_bytes() == b"keep existing export"
    assert set(ALL_TEMP_FILES) == temp_files_before
    assert not list(tmp_path.glob(".rebate-export-*"))


def test_excel_long_filter_list_is_not_silently_truncated(make_csv, tmp_path):
    dataset = load_csv(make_csv())
    ids = tuple(str(100000000000000000 + i) for i in range(1800))
    target = tmp_path / "too-long.xlsx"
    with pytest.raises(DataError, match="32767.*CSV"):
        export_xlsx(target, dataset, dataset.query(Filters(ids)))
    assert not target.exists()


def test_excel_error_like_label_is_exported_as_literal_text(make_csv, tmp_path):
    dataset = load_csv(make_csv([["#N/A", "001", "0", "1", "2026-08-18"]]))
    target = tmp_path / "literal.xlsx"
    export_xlsx(target, dataset, dataset.query(Filters(order_type="#N/A")))
    workbook = load_workbook(target)
    cell = workbook.active.cell(2, 4)
    assert cell.value == "#N/A"
    assert cell.data_type == "s"
    workbook.close()


def test_excel_row_limit_includes_header(make_csv, tmp_path, monkeypatch):
    dataset = load_csv(make_csv())
    # Use a small limit to exercise both sides without constructing a million rows.
    monkeypatch.setattr("exporter.XLSX_MAX_ROWS", 6)
    with pytest.raises(DataError, match="含表头.*CSV"):
        export_xlsx(tmp_path / "overflow.xlsx", dataset, dataset.query())
    monkeypatch.setattr("exporter.XLSX_MAX_ROWS", 7)
    assert export_xlsx(tmp_path / "fits.xlsx", dataset, dataset.query()) == 6


def test_failed_workbook_save_cleans_worksheet_and_target_temporary_files(make_csv, tmp_path, monkeypatch):
    from openpyxl import Workbook
    from openpyxl.worksheet._writer import ALL_TEMP_FILES
    dataset = load_csv(make_csv())
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"previous export")
    temp_files_before = set(ALL_TEMP_FILES)
    def fail(*args):
        raise OSError("simulated disk full")
    monkeypatch.setattr(Workbook, "save", fail)
    with pytest.raises(OSError, match="disk full"):
        export_xlsx(target, dataset, dataset.query())
    assert target.read_bytes() == b"previous export"
    assert set(ALL_TEMP_FILES) == temp_files_before
    assert not list(tmp_path.glob(".rebate-export-*"))
