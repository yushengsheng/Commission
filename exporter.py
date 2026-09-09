"""Export one aggregation view, never overlapping subtotals in a single amount column."""
from __future__ import annotations

import csv
import os
import re
import tempfile
from contextlib import suppress
from pathlib import Path

from rebate_engine import DataError, Dataset, QueryResult, STATUS_LABELS, format_money

EXPORT_HEADERS = [
    "来源文件", "开始日期(含)", "结束日期(含)", "订单类型筛选", "好友ID筛选",
    "好友ID（现货）", "日期", "返佣笔数", "返佣收入(USDT)",
    "所选文件中存在", "UID状态", "全批次笔数", "视图状态筛选", "视图UID搜索",
]
TEXT_COLUMNS = frozenset({0, 1, 2, 3, 4, 5, 6, 8, 9, 10, 12, 13})
VIEW_TITLES = {"id": "按好友", "day": "按日期", "daily_id": "日期与好友"}
XLSX_MAX_ROWS = 1_048_576
XLSX_MAX_CELL_LENGTH = 32_767
INVALID_XML_TEXT = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def safe_cell(value: str) -> str:
    # Defend spreadsheet formula injection in filenames and user-controlled labels.
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def validate_export_target(path, dataset):
    target = Path(path).expanduser().resolve()
    target_stat = target.stat() if target.exists() else None
    for source in dataset.sources:
        same_identity = (target_stat is not None and source.inode and
                         (target_stat.st_dev, target_stat.st_ino) == (source.device, source.inode))
        if target == source.path or same_identity:
            raise DataError("不能覆盖源 CSV，请另选导出文件名")
        # Also protect a replacement source at the original path or its alias.
        if target_stat is not None and source.path.exists() and os.path.samefile(target, source.path):
            raise DataError("不能覆盖源 CSV，请另选导出文件名")
    return target


def _view_rows(result: QueryResult, view: str, status: str, search: str):
    if view not in VIEW_TITLES:
        raise DataError("未知导出类型")
    if view == "id":
        rows = [(uid, "", result.per_id[uid]) for uid in sorted(result.visible_ids(status, search))]
    elif view == "day":
        rows = [("", day.isoformat(), aggregate) for day, aggregate in sorted(result.per_day.items())]
    else:
        rows = [(uid, day.isoformat(), aggregate)
                for (day, uid), aggregate in sorted(result.daily_ids.items())]
    return rows


def _records(dataset: Dataset, result: QueryResult, view: str, status: str, search: str):
    rows = _view_rows(result, view, status, search)
    filters = result.filters
    sources = " | ".join(source.path.name for source in dataset.sources)
    prefix = [safe_cell(sources), str(filters.start or "不限"), str(filters.end or "不限"),
              safe_cell(filters.order_type or "全部"),
              safe_cell(" ".join(dict.fromkeys(filters.ids)) or "全部")]
    def iterate():
        for index, (uid, day, aggregate) in enumerate(rows):
            # Potentially long provenance/ID lists are written ONCE, not once per result row.
            metadata = prefix if index == 0 else ["", *prefix[1:4], ""]
            uid_status = result.statuses.get(uid, "") if view == "id" else ""
            yield metadata + [uid, day, aggregate.count, format_money(aggregate.total),
                ("否" if uid_status == "missing" else "是") if uid_status else "",
                STATUS_LABELS.get(uid_status, ""), result.lifetime_counts.get(uid, 0) if uid_status else "",
                safe_cell(status) if view == "id" else "", safe_cell(search) if view == "id" else ""]
    return len(rows), iterate


def export_csv(path: str | Path, dataset: Dataset, result: QueryResult, view="id",
               status="", search="") -> int:
    target = validate_export_target(path, dataset)
    count, iter_records = _records(dataset, result, view, status, search)
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8-sig", newline="", dir=target.parent,
                                         prefix=".rebate-export-", suffix=".tmp", delete=False) as stream:
            temp_path = Path(stream.name)
            writer = csv.writer(stream)
            writer.writerow(EXPORT_HEADERS)
            writer.writerows(iter_records())
            stream.flush()
            os.fsync(stream.fileno())
        validate_export_target(target, dataset)
        os.replace(temp_path, target)
        temp_path = None
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)
    return count


def export_xlsx(path: str | Path, dataset: Dataset, result: QueryResult, view="id",
                status="", search="") -> int:
    target = validate_export_target(path, dataset)
    # Excel support must not prevent launching the app or using the CSV workflow.
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise DataError("Excel 导出依赖缺失或损坏。请在项目目录执行："
                        ".venv/bin/python -m pip install -r requirements.txt；也可先导出 CSV。") from exc
    count, iter_records = _records(dataset, result, view, status, search)
    if count + 1 > XLSX_MAX_ROWS:
        raise DataError(f"Excel 单表最多 {XLSX_MAX_ROWS} 行（含表头），请缩小筛选范围或导出 CSV。")
    # Validate before openpyxl can truncate cells or create a partial worksheet.
    for record in iter_records():
        for column, value in enumerate(record):
            if not isinstance(value, str):
                continue
            if len(value) > XLSX_MAX_CELL_LENGTH:
                raise DataError(f"“{EXPORT_HEADERS[column]}”超过 Excel 单元格 {XLSX_MAX_CELL_LENGTH} 字符限制，"
                                "请缩小筛选范围或导出 CSV 以保留完整内容。")
            if INVALID_XML_TEXT.search(value):
                raise DataError(f"“{EXPORT_HEADERS[column]}”含 Excel 不支持的字符，请导出 CSV 以保留原值。")
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet(VIEW_TITLES[view])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="237E64")

    def make_cell(value, column, header=False):
        cell = WriteOnlyCell(worksheet, value=value)
        if isinstance(value, str):
            # Formatting alone does not stop labels such as '#N/A' becoming errors.
            cell.data_type = "s"
        if header:
            cell.font, cell.fill = header_font, header_fill
        elif column in TEXT_COLUMNS:
            cell.number_format = "@"
        return cell

    temp_path = None
    try:
        worksheet.freeze_panes = "A2"
        widths = (28, 13, 13, 18, 28, 20, 13, 12, 24, 16, 24, 14, 18, 20)
        for column, width in enumerate(widths, 1):
            worksheet.column_dimensions[get_column_letter(column)].width = width
        worksheet.append([make_cell(value, column, header=True)
                          for column, value in enumerate(EXPORT_HEADERS)])
        for record in iter_records():
            worksheet.append([make_cell(value, column) for column, value in enumerate(record)])
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_HEADERS))}{count + 1}"
        workbook.properties.title = f"返佣汇总-{VIEW_TITLES[view]}"
        workbook.properties.creator = "返佣结算"
        with tempfile.NamedTemporaryFile(dir=target.parent, prefix=".rebate-export-",
                                         suffix=".xlsx", delete=False) as stream:
            temp_path = Path(stream.name)
        workbook.save(temp_path)
        with temp_path.open("rb") as stream:
            os.fsync(stream.fileno())
        validate_export_target(target, dataset)
        os.replace(temp_path, target)
        temp_path = None
    finally:
        # Workbook.close() alone does not remove write-only worksheet temp files
        # when saving fails. Close the XML stream before removing that owned file.
        try:
            writer = worksheet._writer
            if writer is not None:
                if not worksheet.closed:
                    with suppress(OSError, ValueError):
                        worksheet.close()
                with suppress(OSError, ValueError):
                    writer.close()
                if os.path.exists(writer.out):
                    writer.cleanup()
        finally:
            workbook.close()
            if temp_path is not None:
                temp_path.unlink(missing_ok=True)
    return count


def export_result(path: str | Path, dataset: Dataset, result: QueryResult, view="id",
                  status="", search="") -> int:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        return export_csv(path, dataset, result, view, status, search)
    if suffix == ".xlsx":
        return export_xlsx(path, dataset, result, view, status, search)
    raise DataError("导出文件须为 CSV 或 XLSX 格式")
